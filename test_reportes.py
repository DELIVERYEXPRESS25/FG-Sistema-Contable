#!/usr/bin/env python3
"""Tests de funcionalidad para el módulo de Reportes (exportación Excel)."""
import sys
sys.path.insert(0, ".")

from app import app
from services.store import load_data
from services.calculos import (
    calcular_mayor,
    calcular_balanza,
    calcular_estado_resultados,
    calcular_balance_general_con_utilidad,
)
from io import BytesIO
from openpyxl import load_workbook

PASS = 0
FAIL = 0


def ok(msg):
    global PASS
    PASS += 1
    print(f"  ✅ {msg}")


def fail(msg, detail=""):
    global FAIL
    FAIL += 1
    print(f"  ❌ {msg}")
    if detail:
        print(f"     → {detail}")


def get_cell_values(ws, header_row=5):
    """Lee todas las filas de datos de una hoja, retorna lista de dicts."""
    rows = []
    headers = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        if v:
            headers.append((c, str(v).strip()))
    for r in range(header_row + 1, ws.max_row + 1):
        row_data = {}
        for c, h in headers:
            val = ws.cell(row=r, column=c).value
            if val is not None:
                row_data[h] = val
        if row_data:
            rows.append(row_data)
    return rows


def export(hojas, desde=None, hasta=None):
    """Hace POST a /reportes/exportar y retorna el Workbook."""
    with app.test_client() as client:
        form_data = {"hojas": hojas}
        if desde:
            form_data["desde"] = desde
        if hasta:
            form_data["hasta"] = hasta
        # Login first
        client.post("/login", data={"password": "WIZ2026"}, follow_redirects=True)
        resp = client.post("/reportes/exportar", data=form_data)
        assert resp.status_code == 200, f"HTTP {resp.status_code}: {resp.data[:200]}"
        return load_workbook(BytesIO(resp.data))


def test_kardex_orden_cronologico():
    print("\n📦 KARDEX - Orden Cronológico")
    wb = export(["kardex"])
    ws = wb["Kardex PEPS"]
    r = 6
    while r <= ws.max_row:
        nombre = ws.cell(row=r, column=1).value
        if nombre and ws.cell(row=r + 1, column=2).value:
            fechas = []
            r += 1
            while r <= ws.max_row and ws.cell(row=r, column=2).value:
                val = str(ws.cell(row=r, column=2).value)
                if "-" in val:
                    fechas.append(val)
                r += 1
            if len(fechas) >= 2:
                if fechas == sorted(fechas):
                    ok(f"{nombre}: {len(fechas)} movimientos ordenados")
                else:
                    fail(f"{nombre}: NO está en orden", f"Primeras: {fechas[:3]}")
        else:
            r += 1


def test_kardex_filtro_fechas():
    print("\n📦 KARDEX - Filtro de Fechas (Oct 2025)")
    wb = export(["kardex"], desde="2025-10-01", hasta="2025-10-31")
    ws = wb["Kardex PEPS"]
    r = 6
    fuera_de_rango = 0
    total_movs = 0
    while r <= ws.max_row:
        fecha = ws.cell(row=r, column=2).value
        nombre = ws.cell(row=r, column=1).value
        if fecha and nombre is None:
            fecha_str = str(fecha)
            total_movs += 1
            if fecha_str < "2025-10-01" or fecha_str > "2025-10-31":
                fuera_de_rango += 1
                fail(f"  Fecha fuera de rango: {fecha_str}")
        r += 1
    if fuera_de_rango == 0:
        ok(f"Todos los {total_movs} movimientos dentro de Oct 2025")
    elif fuera_de_rango > 0:
        fail(f"{fuera_de_rango} movimientos fuera de rango de {total_movs}")


def test_kardex_saldos_correctos():
    print("\n📦 KARDEX - Saldos Recalculados")
    wb = export(["kardex"])
    ws = wb["Kardex PEPS"]
    r = 6
    while r <= ws.max_row:
        nombre = ws.cell(row=r, column=1).value
        if nombre and ws.cell(row=r + 1, column=2).value:
            saldo_cant = 0
            ok_producto = True
            r += 1
            while r <= ws.max_row and ws.cell(row=r, column=2).value:
                tipo = ws.cell(row=r, column=3).value
                cant = ws.cell(row=r, column=4).value or 0
                saldo_excel = ws.cell(row=r, column=7).value or 0
                if tipo == "Compra":
                    saldo_cant += cant
                elif tipo in ("Salida", "Venta POS"):
                    saldo_cant -= cant
                if saldo_excel != saldo_cant:
                    ok_producto = False
                    fail(f"{nombre} row {r}: saldo esperado {saldo_cant}, Excel tiene {saldo_excel}")
                    break
                r += 1
            if ok_producto and saldo_cant is not None:
                ok(f"{nombre}: saldos consistentes")
        else:
            r += 1


def test_diario_filtro():
    print("\n📖 LIBRO DIARIO - Filtro de Fechas (Oct 2025)")
    wb = export(["diario"], desde="2025-10-01", hasta="2025-10-31")
    ws = wb["Libro Diario"]
    fechas_fuera = []
    for r in range(6, ws.max_row + 1):
        fecha = ws.cell(row=r, column=2).value
        if fecha and str(fecha).strip():
            if str(fecha) < "2025-10-01" or str(fecha) > "2025-10-31":
                fechas_fuera.append(str(fecha))
    if not fechas_fuera:
        ok("Todas las fechas dentro de Oct 2025")
    else:
        fail(f"{len(fechas_fuera)} fechas fuera de rango", str(fechas_fuera[:5]))


def test_diario_totales_cuadran():
    print("\n📖 LIBRO DIARIO - Debe = Haber")
    wb = export(["diario"])
    ws = wb["Libro Diario"]
    total_row = ws.max_row
    debe = ws.cell(row=total_row, column=7).value or 0
    haber = ws.cell(row=total_row, column=8).value or 0
    if abs(debe - haber) < 0.01:
        ok(f"Debe ({debe:.2f}) = Haber ({haber:.2f})")
    else:
        fail(f"No cuadra: Debe={debe:.2f} vs Haber={haber:.2f}")


def test_mayor_saldo_inicial():
    print("\n📗 LIBRO MAYOR - Saldos Iniciales con filtro")
    wb = export(["mayor"], desde="2025-11-01", hasta="2025-11-30")
    ws = wb["Libro Mayor"]
    tiene_ini = False
    for r in range(6, ws.max_row + 1):
        ini_d = ws.cell(row=r, column=4).value or 0
        ini_h = ws.cell(row=r, column=5).value or 0
        if ini_d > 0 or ini_h > 0:
            tiene_ini = True
            break
    if tiene_ini:
        ok("Al menos una cuenta tiene saldo inicial con filtro desde")
    else:
        fail("Ninguna cuenta tiene saldo inicial con filtro desde")


def test_balanza_filtro():
    print("\n⚖️ BALANZA - Filtro Nov 2025")
    wb = export(["balanza"], desde="2025-11-01", hasta="2025-11-30")
    ws = wb["Balanza"]
    subtitle = ws.cell(row=3, column=1).value or ""
    if "2025-11-01" in str(subtitle) and "2025-11-30" in str(subtitle):
        ok("Subtítulo muestra período correcto")
    else:
        fail(f"Subtítulo no muestra período", str(subtitle))


def test_er_filtro():
    print("\n📊 ESTADO DE RESULTADOS - Filtro Oct 2025")
    wb_er = export(["er"], desde="2025-10-01", hasta="2025-10-31")
    wb_all = export(["er"])

    ws_er = wb_er["Estado de Resultados"]
    ws_all = wb_all["Estado de Resultados"]

    util_er = None
    util_all = None
    for r in range(1, ws_er.max_row + 1):
        if ws_er.cell(row=r, column=1).value == "UTILIDAD NETA":
            util_er = ws_er.cell(row=r, column=2).value
    for r in range(1, ws_all.max_row + 1):
        if ws_all.cell(row=r, column=1).value == "UTILIDAD NETA":
            util_all = ws_all.cell(row=r, column=2).value

    if util_er is not None and util_all is not None:
        if abs(util_er) <= abs(util_all) + 0.01:
            ok(f"Utilidad filtrada ({util_er:.2f}) ≤ acumulada ({util_all:.2f})")
        else:
            fail(f"Utilidad filtrada ({util_er:.2f}) > acumulada ({util_all:.2f})")
    else:
        fail("No se encontró Utilidad Neta")


def test_bg_filtro():
    print("\n🏦 BALANCE GENERAL - Filtro Oct vs Nov")
    wb_oct = export(["bg"], desde="2025-10-01", hasta="2025-10-31")
    wb_nov = export(["bg"], desde="2025-11-01", hasta="2025-11-30")

    def get_activos(wb):
        ws = wb["Balance General"]
        for r in range(1, ws.max_row + 1):
            if ws.cell(row=r, column=1).value == "Total Activos":
                return ws.cell(row=r, column=2).value or 0
        return 0

    def get_pasos_cap(wb):
        ws = wb["Balance General"]
        for r in range(1, ws.max_row + 1):
            val = ws.cell(row=r, column=1).value
            if val == "VERIFICACIÓN: Pasivos + Capital":
                return ws.cell(row=r, column=2).value or 0
        return 0

    a_oct = get_activos(wb_oct)
    pc_oct = get_pasos_cap(wb_oct)
    if abs(a_oct - pc_oct) < 0.01:
        ok(f"Oct: Activos ({a_oct:.2f}) = Pas+Cap ({pc_oct:.2f})")
    else:
        fail(f"Oct no cuadra: Activos={a_oct:.2f}, Pas+Cap={pc_oct:.2f}")

    a_nov = get_activos(wb_nov)
    pc_nov = get_pasos_cap(wb_nov)
    if abs(a_nov - pc_nov) < 0.01:
        ok(f"Nov: Activos ({a_nov:.2f}) = Pas+Cap ({pc_nov:.2f})")
    else:
        fail(f"Nov no cuadra: Activos={a_nov:.2f}, Pas+Cap={pc_nov:.2f}")


def test_pos_filtro():
    print("\n🛒 POS - Filtro Dic 2025")
    wb = export(["pos"], desde="2025-12-01", hasta="2025-12-31")
    ws = wb["Ventas POS"]
    fechas_fuera = []
    for r in range(6, ws.max_row + 1):
        fecha = ws.cell(row=r, column=2).value
        if fecha and str(fecha).strip() and str(fecha).strip() != "":
            if str(fecha) < "2025-12-01" or str(fecha) > "2025-12-31":
                fechas_fuera.append(str(fecha))
    if not fechas_fuera:
        ok("Todas las ventas POS dentro de Dic 2025")
    else:
        fail(f"{len(fechas_fuera)} ventas fuera de rango", str(fechas_fuera[:5]))


def test_pos_totales():
    print("\n🛒 POS - Totales")
    wb = export(["pos"])
    ws = wb["Ventas POS"]
    total_row = ws.max_row
    total_excel = ws.cell(row=total_row, column=6).value or 0
    sum_ventas = 0
    for r in range(6, total_row):
        v = ws.cell(row=r, column=6).value
        if v:
            sum_ventas += v
    if abs(total_excel - sum_ventas) < 0.01:
        ok(f"Total ({total_excel:.2f}) = suma ({sum_ventas:.2f})")
    else:
        fail(f"Total no cuadra: {total_excel:.2f} vs {sum_ventas:.2f}")


def test_solo_una_hoja():
    print("\n📄 EXPORTACIÓN SELECTIVA - Solo Kardex")
    wb = export(["kardex"])
    hojas = wb.sheetnames
    if "Kardex PEPS" in hojas and len(hojas) == 1:
        ok("Solo se generó la hoja Kardex PEPS")
    else:
        fail(f"Se generaron hojas inesperadas: {hojas}")


def test_solo_hoja_diario():
    print("\n📄 EXPORTACIÓN SELECTIVA - Solo Diario")
    wb = export(["diario"])
    hojas = wb.sheetnames
    if "Libro Diario" in hojas and len(hojas) == 1:
        ok("Solo se generó la hoja Libro Diario")
    else:
        fail(f"Se generaron hojas inesperadas: {hojas}")


def test_sin_filtro():
    print("\n📅 SIN FILTRO - Acumulado")
    wb = export(["diario", "mayor", "balanza", "er", "bg", "pos", "kardex"])
    hojas = wb.sheetnames
    esperadas = {"Libro Diario", "Libro Mayor", "Balanza", "Estado de Resultados", "Balance General", "Ventas POS", "Kardex PEPS"}
    if esperadas.issubset(set(hojas)):
        ok(f"Se generaron {len(hojas)} hojas correctamente")
    else:
        faltan = esperadas - set(hojas)
        fail(f"Faltan hojas: {faltan}")


def test_kardex_titulo_periodo():
    print("\n📦 KARDEX - Título con Período")
    wb = export(["kardex"], desde="2025-10-01", hasta="2025-12-31")
    ws = wb["Kardex PEPS"]
    sub = ws.cell(row=3, column=1).value or ""
    if "2025-10-01" in str(sub) and "2025-12-31" in str(sub):
        ok("Título muestra período correcto")
    else:
        fail(f"Título incorrecto: {sub}")


if __name__ == "__main__":
    print("=" * 60)
    print("🧪 TESTS DE FUNCIONALIDAD - MÓDULO DE REPORTES")
    print("=" * 60)

    with app.app_context():
        test_kardex_orden_cronologico()
        test_kardex_filtro_fechas()
        test_kardex_saldos_correctos()
        test_kardex_titulo_periodo()
        test_diario_filtro()
        test_diario_totales_cuadran()
        test_mayor_saldo_inicial()
        test_balanza_filtro()
        test_er_filtro()
        test_bg_filtro()
        test_pos_filtro()
        test_pos_totales()
        test_solo_una_hoja()
        test_solo_hoja_diario()
        test_solo_hoja_diario()
        test_sin_filtro()

    print("\n" + "=" * 60)
    total = PASS + FAIL
    print(f"📊 RESULTADO: {PASS}/{total} pasaron, {FAIL} fallaron")
    if FAIL == 0:
        print("🎉 ¡TODOS LOS TESTS PASARON!")
    else:
        print("⚠️  HAY TESTS QUE FALLARON")
    print("=" * 60)
    sys.exit(1 if FAIL else 0)
