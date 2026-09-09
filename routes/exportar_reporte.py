from flask import Blueprint, request, send_file
from services.store import load_data
from services.calculos import (
    calcular_mayor,
    calcular_balanza,
    calcular_estado_resultados,
    calcular_balance_general_con_utilidad,
)
from services.helpers import tipo_saldo

exportar_reporte_bp = Blueprint("exportar_reporte", __name__)


def _parse_fecha_excel(f):
    from datetime import date
    try:
        return date.fromisoformat(f)
    except (ValueError, TypeError):
        return date.today()


@exportar_reporte_bp.route("/reportes/exportar", methods=["POST"], endpoint="exportar_reporte")
def exportar_reporte():
    """Genera un .xlsx con los reportes solicitados (envío directo, no guarda en disco)."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        import subprocess, sys

        subprocess.check_call(
            [sys.executable, "-m", "pip", "install", "openpyxl", "--quiet"]
        )
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

    data = load_data()
    desde = request.form.get("desde", "") or None
    hasta = request.form.get("hasta", "") or None
    tipo_cuenta = request.form.get("tipo_cuenta", "") or None
    cuenta_especifica = request.form.get("cuenta_especifica", "") or None
    from datetime import date, datetime

    hojas_raw = request.form.getlist("hojas")
    if hojas_raw:
        hojas_seleccionadas = set(hojas_raw)
    else:
        tipo_reporte = request.form.get("tipo", "todos")
        if tipo_reporte == "todos":
            hojas_seleccionadas = {"diario", "mayor", "balanza", "er", "bg", "pos", "kardex", "cobros"}
        else:
            mapa_legacy = {
                "diario": "diario", "mayor": "mayor", "balanza": "balanza",
                "estado_resultados": "er", "balance_general": "bg",
                "kardex": "kardex", "antiguedad_cobros": "cobros",
            }
            hojas_seleccionadas = {mapa_legacy.get(tipo_reporte, tipo_reporte)}

    def _filtrar_cuentas(movimientos):
        """Filtra movimientos por tipo de cuenta o cuenta específica."""
        if not tipo_cuenta and not cuenta_especifica:
            return movimientos
        filtered = []
        for mov in movimientos:
            cod = mov.get("cuenta", "")
            info = cuentas.get(cod, {})
            if cuenta_especifica and cod != cuenta_especifica:
                continue
            if tipo_cuenta and info.get("tipo", "") != tipo_cuenta:
                continue
            filtered.append(mov)
        return filtered

    data_original = data  # Guardar referencia a datos completos para saldos iniciales
    if desde or hasta:
        import copy
        data = copy.deepcopy(data)
        if desde and hasta:
            data["diario"] = [e for e in data.get("diario", []) if desde <= e.get("fecha", "") <= hasta]
            data["ajustes"] = [a for a in data.get("ajustes", []) if desde <= a.get("fecha", "") <= hasta]
        elif desde:
            data["diario"] = [e for e in data.get("diario", []) if desde <= e.get("fecha", "")]
            data["ajustes"] = [a for a in data.get("ajustes", []) if desde <= a.get("fecha", "")]
        elif hasta:
            data["diario"] = [e for e in data.get("diario", []) if e.get("fecha", "") <= hasta]
            data["ajustes"] = [a for a in data.get("ajustes", []) if a.get("fecha", "") <= hasta]

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1A1D2E")
    accent_fill = PatternFill("solid", fgColor="4F8CFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1A1D2E")
    sub_font = Font(name="Calibri", size=10, color="7A7F99")
    money_fmt = "#,##0.00"
    thin_border = Border(
        left=Side(style="thin", color="D0D3E0"),
        right=Side(style="thin", color="D0D3E0"),
        top=Side(style="thin", color="D0D3E0"),
        bottom=Side(style="thin", color="D0D3E0"),
    )
    total_fill = PatternFill("solid", fgColor="E8EAEF")
    total_font = Font(name="Calibri", size=11, bold=True, color="1A1D2E")
    green_font = Font(name="Calibri", size=11, color="27AE60")
    red_font = Font(name="Calibri", size=11, color="E74C3C")

    def style_header_row(ws, row, cols):
        for c in range(1, cols + 1):
            cell = ws.cell(row=row, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

    def style_data_row(ws, row, cols):
        for c in range(1, cols + 1):
            cell = ws.cell(row=row, column=c)
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10, color="333333")

    def style_total_row(ws, row, cols):
        for c in range(1, cols + 1):
            cell = ws.cell(row=row, column=c)
            cell.font = total_font
            cell.fill = total_fill
            cell.border = Border(
                left=Side(style="thin", color="D0D3E0"),
                right=Side(style="thin", color="D0D3E0"),
                top=Side(style="medium", color="4F8CFF"),
                bottom=Side(style="thin", color="D0D3E0"),
            )

    def add_title(ws, title, subtitle=""):
        ws.cell(row=1, column=1, value="F & G — Sistema Contable").font = Font(
            name="Calibri", size=12, bold=True, color="4F8CFF"
        )
        ws.cell(row=2, column=1, value=title).font = title_font
        if subtitle:
            ws.cell(row=3, column=1, value=subtitle).font = sub_font
        ws.cell(
            row=3 if not subtitle else 4,
            column=1,
            value=f"Fecha de exportación: {date.today().isoformat()}",
        ).font = sub_font
        return 5  # primera fila de datos

    wb = Workbook()
    wb.remove(wb.active)  # borrar hoja por defecto

    mayor = calcular_mayor(data)
    cuentas = data["cuentas"]

    if "diario" in hojas_seleccionadas:
        periodo_str = f"Período: {desde} al {hasta}" if desde and hasta else "Acumulado"
        ws = wb.create_sheet("Libro Diario")
        ws.sheet_properties.tabColor = "4F8CFF"
        start = add_title(ws, "Libro Diario", f"Registro cronológico de asientos — {periodo_str}")
        headers = [
            "#",
            "Fecha",
            "Descripción",
            "Ref",
            "Cuenta",
            "Nombre Cuenta",
            "Debe C$",
            "Haber C$",
        ]
        for i, h in enumerate(headers, 1):
            ws.cell(row=start, column=i, value=h)
        style_header_row(ws, start, len(headers))
        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 14
        ws.column_dimensions["C"].width = 38
        ws.column_dimensions["D"].width = 10
        ws.column_dimensions["E"].width = 10
        ws.column_dimensions["F"].width = 28
        ws.column_dimensions["G"].width = 16
        ws.column_dimensions["H"].width = 16

        r = start + 1
        total_debe = total_haber = 0
        all_entries = list(data["diario"]) + list(data.get("ajustes", []))
        all_entries.sort(key=lambda x: x.get("fecha", ""))
        for entry in all_entries:
            movs_filtrados = _filtrar_cuentas(entry["movimientos"])
            if not movs_filtrados:
                continue
            for idx, mov in enumerate(movs_filtrados):
                ws.cell(row=r, column=1, value=entry["id"] if idx == 0 else "")
                ws.cell(row=r, column=2, value=entry["fecha"] if idx == 0 else "")
                ws.cell(row=r, column=3, value=entry["descripcion"] if idx == 0 else "")
                ws.cell(row=r, column=4, value=entry.get("ref", "AJ") if idx == 0 else "")
                ws.cell(row=r, column=5, value=mov["cuenta"])
                ws.cell(
                    row=r,
                    column=6,
                    value=cuentas.get(mov["cuenta"], {}).get("nombre", mov["cuenta"]),
                )
                debe = mov["monto"] if mov["tipo"] == "Debe" else 0
                haber = mov["monto"] if mov["tipo"] == "Haber" else 0
                ws.cell(row=r, column=7, value=debe).number_format = money_fmt
                ws.cell(row=r, column=8, value=haber).number_format = money_fmt
                if debe:
                    ws.cell(row=r, column=7).font = green_font
                if haber:
                    ws.cell(row=r, column=8).font = red_font
                total_debe += debe
                total_haber += haber
                style_data_row(ws, r, len(headers))
                r += 1
        ws.cell(row=r, column=6, value="TOTALES")
        ws.cell(row=r, column=7, value=total_debe).number_format = money_fmt
        ws.cell(row=r, column=8, value=total_haber).number_format = money_fmt
        style_total_row(ws, r, len(headers))

    if "mayor" in hojas_seleccionadas:
        periodo_str = f"Período: {desde} al {hasta}" if desde and hasta else "Acumulado"
        ws = wb.create_sheet("Libro Mayor")
        ws.sheet_properties.tabColor = "7C5CFC"
        start = add_title(ws, "Libro Mayor", f"Cuentas T — Movimientos por cuenta — {periodo_str}")

        if desde:
            data_ini = {k: v for k, v in data_original.items()}
            data_ini["diario"] = [e for e in data_original.get("diario", []) if e.get("fecha", "") < desde]
            data_ini["ajustes"] = [a for a in data_original.get("ajustes", []) if a.get("fecha", "") < desde]
            mayor_ini = calcular_mayor(data_ini)
        else:
            mayor_ini = None

        headers = ["Cuenta", "Nombre", "Tipo", "Saldo Inicial Deudor", "Saldo Inicial Acreedor", "Debe Período", "Haber Período", "Saldo Final Deudor", "Saldo Final Acreedor"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=start, column=i, value=h)
        style_header_row(ws, start, len(headers))
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 12
        ws.column_dimensions["D"].width = 16
        ws.column_dimensions["E"].width = 16
        ws.column_dimensions["F"].width = 16
        ws.column_dimensions["G"].width = 16
        ws.column_dimensions["H"].width = 16
        ws.column_dimensions["I"].width = 16

        r = start + 1
        for codigo in sorted(cuentas.keys()):
            info = cuentas[codigo]
            if cuenta_especifica and codigo != cuenta_especifica:
                continue
            if tipo_cuenta and info.get("tipo", "") != tipo_cuenta:
                continue
            debe_mov = mayor[codigo]["debe"] if codigo in mayor else 0
            haber_mov = mayor[codigo]["haber"] if codigo in mayor else 0
            ts = tipo_saldo(info["tipo"])

            if mayor_ini is not None and codigo in mayor_ini:
                ini_d = mayor_ini[codigo]["debe"]
                ini_h = mayor_ini[codigo]["haber"]
            else:
                ini_d = 0
                ini_h = 0

            if ts == "Debe":
                saldo_ini_d = max(ini_d - ini_h, 0)
                saldo_ini_h = max(ini_h - ini_d, 0)
                saldo_fin_d = max((ini_d - ini_h) + (debe_mov - haber_mov), 0)
                saldo_fin_h = max(-(ini_d - ini_h) - (debe_mov - haber_mov), 0)
            else:
                saldo_ini_d = max(-(ini_h - ini_d), 0)
                saldo_ini_h = max(ini_h - ini_d, 0)
                saldo_fin_d = max(-(ini_h - ini_d) - (haber_mov - debe_mov), 0)
                saldo_fin_h = max((ini_h - ini_d) + (haber_mov - debe_mov), 0)

            ws.cell(row=r, column=1, value=codigo)
            ws.cell(row=r, column=2, value=info["nombre"])
            ws.cell(row=r, column=3, value=info["tipo"])
            ws.cell(row=r, column=4, value=saldo_ini_d).number_format = money_fmt
            ws.cell(row=r, column=5, value=saldo_ini_h).number_format = money_fmt
            ws.cell(row=r, column=6, value=debe_mov).number_format = money_fmt
            ws.cell(row=r, column=7, value=haber_mov).number_format = money_fmt
            ws.cell(row=r, column=8, value=saldo_fin_d).number_format = money_fmt
            ws.cell(row=r, column=9, value=saldo_fin_h).number_format = money_fmt
            style_data_row(ws, r, len(headers))
            r += 1

    if "balanza" in hojas_seleccionadas:
        periodo_str = f"Período: {desde} al {hasta}" if desde and hasta else "Acumulado"
        ws = wb.create_sheet("Balanza")
        ws.sheet_properties.tabColor = "F39C12"
        start = add_title(
            ws, "Balanza de Comprobación", f"Verificación de débitos y créditos — {periodo_str}"
        )

        if desde:
            import copy as _copy
            data_completa = _copy.deepcopy(data_original)
            data_completa["diario"] = [e for e in data_original.get("diario", []) if e.get("fecha", "") < desde]
            data_completa["ajustes"] = [a for a in data_original.get("ajustes", []) if a.get("fecha", "") < desde]
            mayor_inicial = calcular_mayor(data_completa)
        else:
            mayor_inicial = None

        balanza_data, td, th, tsd, tsh = calcular_balanza(data)

        headers = [
            "Código",
            "Cuenta",
            "Tipo",
            "Saldo Inicial Deudor",
            "Saldo Inicial Acreedor",
            "Mov. Deudor",
            "Mov. Acreedor",
            "Saldo Final Deudor",
            "Saldo Final Acreedor",
        ]
        for i, h in enumerate(headers, 1):
            ws.cell(row=start, column=i, value=h)
        style_header_row(ws, start, len(headers))
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 12
        ws.column_dimensions["D"].width = 18
        ws.column_dimensions["E"].width = 18
        ws.column_dimensions["F"].width = 18
        ws.column_dimensions["G"].width = 18
        ws.column_dimensions["H"].width = 18
        ws.column_dimensions["I"].width = 18

        r = start + 1
        tot_sd_i = tot_sh_i = tot_d = tot_h = tot_sd_f = tot_sh_f = 0
        for item in balanza_data:
            if item.get("es_header"):
                ws.cell(row=r, column=1, value=item["codigo"])
                ws.cell(row=r, column=2, value=item["nombre"]).font = Font(
                    name="Calibri", size=10, bold=True, color="1A1D2E"
                )
                style_data_row(ws, r, 9)
                r += 1
                continue
            if item.get("es_subtotal"):
                continue
            codigo = item["codigo"]
            info = cuentas.get(codigo, {})
            tipo = info.get("tipo", "")
            ts = tipo_saldo(tipo)

            debe_mov = item["debe"]
            haber_mov = item["haber"]

            if mayor_inicial is not None and codigo in mayor_inicial:
                ini_d = mayor_inicial[codigo]["debe"]
                ini_h = mayor_inicial[codigo]["haber"]
            else:
                ini_d = 0
                ini_h = 0

            if ts == "Debe":
                saldo_ini_d = max(ini_d - ini_h, 0)
                saldo_ini_h = max(ini_h - ini_d, 0)
                saldo_fin_d = max(saldo_ini_d + saldo_ini_h + (debe_mov - haber_mov), 0) if saldo_ini_d > 0 else max(debe_mov - haber_mov, 0)
                saldo_fin_h = max(saldo_ini_h + saldo_ini_d + (haber_mov - debe_mov), 0) if saldo_ini_h > 0 else max(haber_mov - debe_mov, 0)
                neto_ini = ini_d - ini_h
                neto_fin = neto_ini + (debe_mov - haber_mov)
                saldo_ini_d = max(neto_ini, 0)
                saldo_ini_h = max(-neto_ini, 0)
                saldo_fin_d = max(neto_fin, 0)
                saldo_fin_h = max(-neto_fin, 0)
            else:
                neto_ini = ini_h - ini_d
                neto_fin = neto_ini + (haber_mov - debe_mov)
                saldo_ini_d = max(-neto_ini, 0)
                saldo_ini_h = max(neto_ini, 0)
                saldo_fin_d = max(-neto_fin, 0)
                saldo_fin_h = max(neto_fin, 0)

            ws.cell(row=r, column=1, value=codigo)
            ws.cell(row=r, column=2, value=item["nombre"])
            ws.cell(row=r, column=3, value=tipo)
            ws.cell(row=r, column=4, value=saldo_ini_d).number_format = money_fmt
            ws.cell(row=r, column=5, value=saldo_ini_h).number_format = money_fmt
            ws.cell(row=r, column=6, value=debe_mov).number_format = money_fmt
            ws.cell(row=r, column=7, value=haber_mov).number_format = money_fmt
            ws.cell(row=r, column=8, value=saldo_fin_d).number_format = money_fmt
            ws.cell(row=r, column=9, value=saldo_fin_h).number_format = money_fmt
            style_data_row(ws, r, 9)
            tot_sd_i += saldo_ini_d
            tot_sh_i += saldo_ini_h
            tot_d += debe_mov
            tot_h += haber_mov
            tot_sd_f += saldo_fin_d
            tot_sh_f += saldo_fin_h
            r += 1
        ws.cell(row=r, column=2, value="TOTALES")
        ws.cell(row=r, column=4, value=tot_sd_i).number_format = money_fmt
        ws.cell(row=r, column=5, value=tot_sh_i).number_format = money_fmt
        ws.cell(row=r, column=6, value=tot_d).number_format = money_fmt
        ws.cell(row=r, column=7, value=tot_h).number_format = money_fmt
        ws.cell(row=r, column=8, value=tot_sd_f).number_format = money_fmt
        ws.cell(row=r, column=9, value=tot_sh_f).number_format = money_fmt
        style_total_row(ws, r, 9)

    if "er" in hojas_seleccionadas:
        periodo_str = f"Período: {desde} al {hasta}" if desde and hasta else "Acumulado"
        ws = wb.create_sheet("Estado de Resultados")
        ws.sheet_properties.tabColor = "2ECC71"
        start = add_title(
            ws, "Estado de Resultados", f"Ingresos, Gastos y Utilidad Neta — {periodo_str}"
        )
        di, ti, dg, tg, util = calcular_estado_resultados(data)

        r = start
        ws.cell(row=r, column=1, value="INGRESOS").font = Font(
            name="Calibri", size=11, bold=True, color="27AE60"
        )
        r += 1
        for item in di:
            if item.get("tipo") in ("header", "subtotal"):
                continue
            ws.cell(row=r, column=1, value=item["nombre"]).font = Font(
                name="Calibri", size=10, color="333333"
            )
            ws.cell(row=r, column=2, value=item["saldo"]).number_format = money_fmt
            ws.cell(row=r, column=2).font = green_font
            r += 1
        ws.cell(row=r, column=1, value="Total Ingresos").font = total_font
        ws.cell(row=r, column=2, value=ti).number_format = money_fmt
        ws.cell(row=r, column=2).font = Font(
            name="Calibri", size=11, bold=True, color="27AE60"
        )
        style_total_row(ws, r, 2)
        r += 2

        ws.cell(row=r, column=1, value="GASTOS").font = Font(
            name="Calibri", size=11, bold=True, color="E74C3C"
        )
        r += 1
        for item in dg:
            if item.get("tipo") in ("header", "subtotal"):
                continue
            ws.cell(row=r, column=1, value=item["nombre"]).font = Font(
                name="Calibri", size=10, color="333333"
            )
            ws.cell(row=r, column=2, value=item["saldo"]).number_format = money_fmt
            ws.cell(row=r, column=2).font = red_font
            r += 1
        ws.cell(row=r, column=1, value="Total Gastos").font = total_font
        ws.cell(row=r, column=2, value=tg).number_format = money_fmt
        ws.cell(row=r, column=2).font = Font(
            name="Calibri", size=11, bold=True, color="E74C3C"
        )
        style_total_row(ws, r, 2)
        r += 2

        ws.cell(row=r, column=1, value="UTILIDAD NETA").font = Font(
            name="Calibri", size=13, bold=True, color="1A1D2E"
        )
        ws.cell(row=r, column=2, value=util).number_format = money_fmt
        ws.cell(row=r, column=2).font = Font(
            name="Calibri",
            size=13,
            bold=True,
            color="27AE60" if util >= 0 else "E74C3C",
        )
        style_total_row(ws, r, 2)

        ws.column_dimensions["A"].width = 32
        ws.column_dimensions["B"].width = 20

    if "bg" in hojas_seleccionadas:
        periodo_str = f"Período: {desde} al {hasta}" if desde and hasta else "Acumulado"
        ws = wb.create_sheet("Balance General")
        ws.sheet_properties.tabColor = "E74C3C"
        start = add_title(ws, "Balance General", f"Activos = Pasivos + Capital — {periodo_str}")
        activos, ta, pasivos, tp, capital_items, tc = calcular_balance_general_con_utilidad(data_original, desde, hasta)

        r = start
        ws.cell(row=r, column=1, value="ACTIVOS").font = Font(
            name="Calibri", size=11, bold=True, color="4F8CFF"
        )
        r += 1
        for item in activos:
            if item.get("tipo") == "header":
                ws.cell(row=r, column=1, value=item["nombre"]).font = Font(
                    name="Calibri", size=10, bold=True, color="4F8CFF"
                )
                r += 1
                continue
            if item.get("tipo") == "subtotal":
                ws.cell(row=r, column=1, value="  Total " + item["nombre"]).font = total_font
                ws.cell(row=r, column=2, value=item["saldo"]).number_format = money_fmt
                style_data_row(ws, r, 2)
                r += 1
                continue
            ws.cell(row=r, column=1, value=item["nombre"]).font = Font(
                name="Calibri", size=10, color="333333"
            )
            ws.cell(row=r, column=2, value=item["saldo"]).number_format = money_fmt
            ws.cell(row=r, column=2).font = (
                green_font if item["saldo"] >= 0 else red_font
            )
            r += 1
        ws.cell(row=r, column=1, value="Total Activos").font = total_font
        ws.cell(row=r, column=2, value=ta).number_format = money_fmt
        ws.cell(row=r, column=2).font = Font(
            name="Calibri", size=11, bold=True, color="4F8CFF"
        )
        style_total_row(ws, r, 2)
        r += 2

        ws.cell(row=r, column=1, value="PASIVOS").font = Font(
            name="Calibri", size=11, bold=True, color="E74C3C"
        )
        r += 1
        for item in pasivos:
            if item.get("tipo") == "header":
                ws.cell(row=r, column=1, value=item["nombre"]).font = Font(
                    name="Calibri", size=10, bold=True, color="E74C3C"
                )
                r += 1
                continue
            if item.get("tipo") == "subtotal":
                ws.cell(row=r, column=1, value="  Total " + item["nombre"]).font = total_font
                ws.cell(row=r, column=2, value=item["saldo"]).number_format = money_fmt
                style_data_row(ws, r, 2)
                r += 1
                continue
            ws.cell(row=r, column=1, value=item["nombre"]).font = Font(
                name="Calibri", size=10, color="333333"
            )
            ws.cell(row=r, column=2, value=item["saldo"]).number_format = money_fmt
            ws.cell(row=r, column=2).font = red_font
            r += 1
        ws.cell(row=r, column=1, value="Total Pasivos").font = total_font
        ws.cell(row=r, column=2, value=tp).number_format = money_fmt
        ws.cell(row=r, column=2).font = Font(
            name="Calibri", size=11, bold=True, color="E74C3C"
        )
        style_total_row(ws, r, 2)
        r += 2

        ws.cell(row=r, column=1, value="CAPITAL").font = Font(
            name="Calibri", size=11, bold=True, color="2ECC71"
        )
        r += 1
        for item in capital_items:
            if item.get("tipo") == "header":
                ws.cell(row=r, column=1, value=item["nombre"]).font = Font(
                    name="Calibri", size=10, bold=True, color="2ECC71"
                )
                r += 1
                continue
            if item.get("tipo") == "subtotal":
                ws.cell(row=r, column=1, value="  Total " + item["nombre"]).font = total_font
                ws.cell(row=r, column=2, value=item["saldo"]).number_format = money_fmt
                style_data_row(ws, r, 2)
                r += 1
                continue
            ws.cell(row=r, column=1, value=item["nombre"]).font = Font(
                name="Calibri", size=10, color="333333"
            )
            ws.cell(row=r, column=2, value=item["saldo"]).number_format = money_fmt
            ws.cell(row=r, column=2).font = (
                green_font if item["saldo"] >= 0 else red_font
            )
            r += 1
        ws.cell(row=r, column=1, value="Total Capital").font = total_font
        ws.cell(row=r, column=2, value=tc).number_format = money_fmt
        ws.cell(row=r, column=2).font = Font(
            name="Calibri", size=11, bold=True, color="2ECC71"
        )
        style_total_row(ws, r, 2)
        r += 2

        ws.cell(row=r, column=1, value="VERIFICACIÓN: Pasivos + Capital").font = Font(
            name="Calibri", size=11, bold=True, color="333333"
        )
        ws.cell(row=r, column=2, value=tp + tc).number_format = money_fmt
        ws.cell(row=r, column=2).font = Font(
            name="Calibri", size=11, bold=True, color="4F8CFF"
        )

        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 20

    if "pos" in hojas_seleccionadas:
        pos_historial = data.get("pos_historial", [])
        if desde and hasta:
            pos_historial = [v for v in pos_historial if desde <= v.get("fecha", "") <= hasta]
        elif desde:
            pos_historial = [v for v in pos_historial if desde <= v.get("fecha", "")]
        elif hasta:
            pos_historial = [v for v in pos_historial if v.get("fecha", "") <= hasta]
        if pos_historial:
            periodo_str = f"Período: {desde} al {hasta}" if desde and hasta else "Acumulado"
            ws = wb.create_sheet("Ventas POS")
            ws.sheet_properties.tabColor = "F39C12"
            start = add_title(
                ws, "Ventas POS", f"Historial de ventas — {periodo_str}"
            )
            headers = [
                "Ref",
                "Fecha",
                "Cliente",
                "Forma Pago",
                "Productos",
                "Total C$",
                "Costo C$",
                "Utilidad C$",
            ]
            for i, h in enumerate(headers, 1):
                ws.cell(row=start, column=i, value=h)
            style_header_row(ws, start, len(headers))
            ws.column_dimensions["A"].width = 12
            ws.column_dimensions["B"].width = 14
            ws.column_dimensions["C"].width = 22
            ws.column_dimensions["D"].width = 14
            ws.column_dimensions["E"].width = 40
            ws.column_dimensions["F"].width = 16
            ws.column_dimensions["G"].width = 16
            ws.column_dimensions["H"].width = 16
            r = start + 1
            total_ventas = total_costos = total_utilidades = 0
            for venta in pos_historial:
                productos_str = ", ".join(
                    f"{l.get('nombre', l.get('producto', '?'))} x{l.get('cantidad', 1)}"
                    for l in venta.get("lineas", [])
                )
                total_v = venta.get("total", venta.get("total_venta", 0))
                total_c = venta.get("costo", 0)
                utilidad = venta.get("utilidad", total_v - total_c)
                ws.cell(row=r, column=1, value=venta.get("ref", ""))
                ws.cell(row=r, column=2, value=venta.get("fecha", ""))
                ws.cell(row=r, column=3, value=venta.get("cliente", ""))
                ws.cell(row=r, column=4, value=venta.get("forma_pago", ""))
                ws.cell(row=r, column=5, value=productos_str)
                ws.cell(row=r, column=6, value=total_v).number_format = money_fmt
                ws.cell(row=r, column=6).font = green_font
                ws.cell(row=r, column=7, value=total_c).number_format = money_fmt
                ws.cell(row=r, column=8, value=utilidad).number_format = money_fmt
                ws.cell(row=r, column=8).font = (
                    green_font if utilidad >= 0 else red_font
                )
                style_data_row(ws, r, len(headers))
                total_ventas += total_v
                total_costos += total_c
                total_utilidades += utilidad
                r += 1
            ws.cell(row=r, column=5, value="TOTALES")
            ws.cell(row=r, column=6, value=total_ventas).number_format = money_fmt
            ws.cell(row=r, column=7, value=total_costos).number_format = money_fmt
            ws.cell(row=r, column=8, value=total_utilidades).number_format = money_fmt
            style_total_row(ws, r, len(headers))

    if "kardex" in hojas_seleccionadas and "kardex" in data and data["kardex"]:
        periodo_str = f"Período: {desde} al {hasta}" if desde and hasta else "Acumulado"
        ws = wb.create_sheet("Kardex PEPS")
        ws.sheet_properties.tabColor = "9B59B6"
        start = add_title(ws, "Kardex PEPS", f"Inventario por producto — Método PEPS — {periodo_str}")
        headers = ["Producto", "Fecha", "Tipo Movimiento", "Cantidad", "Costo Unit.", "Costo Total", "Saldo Cant.", "Saldo Costo", "Descripción"]
        widths = [28, 14, 18, 12, 14, 14, 12, 14, 35]
        for i, (h, w) in enumerate(zip(headers, widths), 1):
            ws.cell(row=start, column=i, value=h)
        style_header_row(ws, start, len(headers))
        for ci, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(ci)].width = w
        r = start + 1
        total_saldo_costo = 0
        total_saldo_cant = 0
        num_productos = 0
        productos = sorted(data["kardex"].keys())
        for nombre in productos:
            movimientos = data["kardex"][nombre]
            if not movimientos:
                continue
            movs_filtrados = []
            for mov in movimientos:
                mov_fecha = mov.get("fecha", "")
                if desde and mov_fecha < desde:
                    continue
                if hasta and mov_fecha > hasta:
                    continue
                movs_filtrados.append(mov)
            if not movs_filtrados:
                continue
            num_productos += 1
            ultimo = movs_filtrados[-1]
            ultimo_saldo_cant = ultimo.get("saldo", 0)
            ultimo_saldo_costo = ultimo.get("costo", 0) * ultimo_saldo_cant if ultimo_saldo_cant else 0
            total_saldo_cant += ultimo_saldo_cant
            total_saldo_costo += ultimo_saldo_costo

            ws.cell(row=r, column=1, value=nombre).font = Font(name="Calibri", size=10, bold=True, color="333333")
            r += 1
            for mov in movs_filtrados:
                tipo_mov = mov.get("tipo", "")
                desc = mov.get("descripcion", "")
                if "Venta POS" in desc:
                    tipo_label = "Venta POS"
                elif tipo_mov == "entrada":
                    tipo_label = "Compra"
                elif tipo_mov == "salida":
                    tipo_label = "Salida"
                elif "ajuste" in desc.lower() or "Ajuste" in desc:
                    tipo_label = "Ajuste"
                else:
                    tipo_label = tipo_mov.title() if tipo_mov else "Otro"
                cantidad = mov.get("cantidad", 0)
                costo_unit = mov.get("costo", 0)
                costo_total = mov.get("total", cantidad * costo_unit)
                saldo_cant = mov.get("saldo", 0)
                saldo_costo = mov.get("costo", 0) * saldo_cant if saldo_cant else 0

                ws.cell(row=r, column=2, value=mov.get("fecha", ""))
                ws.cell(row=r, column=3, value=tipo_label)
                ws.cell(row=r, column=4, value=cantidad)
                ws.cell(row=r, column=5, value=costo_unit).number_format = money_fmt
                ws.cell(row=r, column=6, value=costo_total).number_format = money_fmt
                ws.cell(row=r, column=7, value=saldo_cant)
                ws.cell(row=r, column=8, value=saldo_costo).number_format = money_fmt
                ws.cell(row=r, column=9, value=desc)
                if tipo_mov == "salida":
                    for c in range(1, 10):
                        ws.cell(row=r, column=c).font = red_font
                elif tipo_mov == "entrada":
                    for c in range(1, 10):
                        ws.cell(row=r, column=c).font = green_font
                style_data_row(ws, r, len(headers))
                r += 1
            r += 1

        r += 1
        ws.cell(row=r, column=1, value="RESUMEN INVENTARIO").font = Font(
            name="Calibri", size=11, bold=True, color="1A1D2E"
        )
        r += 1
        ws.cell(row=r, column=1, value="Total Productos:")
        ws.cell(row=r, column=2, value=num_productos).font = Font(
            name="Calibri", size=11, bold=True, color="4F8CFF"
        )
        r += 1
        ws.cell(row=r, column=1, value="Total Unidades en Stock:")
        ws.cell(row=r, column=2, value=total_saldo_cant).font = Font(
            name="Calibri", size=11, bold=True, color="4F8CFF"
        )
        r += 1
        ws.cell(row=r, column=1, value="Valor Total Inventario:")
        ws.cell(row=r, column=2, value=total_saldo_costo).number_format = money_fmt
        ws.cell(row=r, column=2).font = Font(
            name="Calibri", size=12, bold=True, color="27AE60"
        )

    if "cobros" in hojas_seleccionadas:
        periodo_str = f"Período: {desde} al {hasta}" if desde and hasta else "Acumulado"
        hoy = date.today()
        cobros = data.get("cuentas_cobrar", [])
        activos = [c for c in cobros if c.get("estado", "pendiente") != "pagado"]
        ws = wb.create_sheet("Antiguedad Cobros")
        ws.sheet_properties.tabColor = "E67E22"
        start = add_title(ws, "Antigüedad de Cuentas por Cobrar",
                          f"Saldos pendientes por rango de vencimiento — {periodo_str}")
        headers = ["Rango", "Cantidad", "Total C$"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=start, column=i, value=h)
        style_header_row(ws, start, len(headers))
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 14
        ws.column_dimensions["C"].width = 20
        rangos = [
            ("0-30 días",   lambda f: 0 <= (hoy - _parse_fecha_excel(f)).days <= 30),
            ("31-60 días",  lambda f: 31 <= (hoy - _parse_fecha_excel(f)).days <= 60),
            ("61-90 días",  lambda f: 61 <= (hoy - _parse_fecha_excel(f)).days <= 90),
            ("91+ días",    lambda f: (hoy - _parse_fecha_excel(f)).days > 90),
        ]
        r = start + 1
        gran_total = 0
        for label, cond in rangos:
            items = [c for c in activos if cond(c.get("fecha", ""))]
            if not items:
                continue
            total_rango = sum(
                c["monto"] - sum(pg.get("monto", 0) for pg in c.get("pagos", []))
                for c in items
            )
            gran_total += total_rango
            ws.cell(row=r, column=1, value=label)
            ws.cell(row=r, column=2, value=len(items))
            ws.cell(row=r, column=3, value=total_rango).number_format = money_fmt
            style_data_row(ws, r, len(headers))
            r += 1
        ws.cell(row=r, column=1, value="TOTAL")
        ws.cell(row=r, column=2, value=len(activos))
        ws.cell(row=r, column=3, value=gran_total).number_format = money_fmt
        style_total_row(ws, r, len(headers))
        r += 2
        ws.cell(row=r, column=1, value="Detalle de cuentas pendientes").font = total_font
        r += 1
        dh = ["#", "Fecha", "Descripción", "Monto C$", "Saldo C$", "Días"]
        for i, h in enumerate(dh, 1):
            ws.cell(row=r, column=i, value=h)
        style_header_row(ws, r, len(dh))
        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 14
        ws.column_dimensions["C"].width = 38
        ws.column_dimensions["D"].width = 16
        ws.column_dimensions["E"].width = 16
        ws.column_dimensions["F"].width = 10
        r += 1
        for c in activos:
            saldo = c["monto"] - sum(pg.get("monto", 0) for pg in c.get("pagos", []))
            dias = (hoy - _parse_fecha_excel(c.get("fecha", ""))).days
            ws.cell(row=r, column=1, value=c.get("id", ""))
            ws.cell(row=r, column=2, value=c.get("fecha", ""))
            ws.cell(row=r, column=3, value=c.get("descripcion", "")[:50])
            ws.cell(row=r, column=4, value=c["monto"]).number_format = money_fmt
            ws.cell(row=r, column=5, value=saldo).number_format = money_fmt
            ws.cell(row=r, column=6, value=dias)
            style_data_row(ws, r, len(dh))
            r += 1

    if not wb.sheetnames:
        ws = wb.create_sheet("Reporte")
        ws.cell(row=1, column=1, value="No hay datos para este reporte.")

    nombre_hojas = {
        "diario": "Diario", "mayor": "Mayor", "balanza": "Balanza",
        "er": "EstadoResultados", "bg": "BalanceGeneral",
        "pos": "VentasPOS", "kardex": "KardexPEPS", "cobros": "AntiguedadCobros",
    }
    if len(hojas_seleccionadas) == 8 or not hojas_seleccionadas:
        label = "Completo"
    elif len(hojas_seleccionadas) == 1:
        label = nombre_hojas.get(list(hojas_seleccionadas)[0], "Reporte")
    else:
        label = "_".join(nombre_hojas.get(h, h) for h in sorted(hojas_seleccionadas))
    fecha_str = date.today().isoformat()
    filename = f"FG_{label}_{fecha_str}.xlsx"

    from io import BytesIO

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )
